"""Relabel the existing paper Figure 14 without changing its layout or media."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
import win32com.client as win32


PAPER = Path(__file__).resolve().parent.parent
SOURCE = PAPER / "figures/q3/fig14_q3_heaf_framework_source.pptx"
PDF = PAPER / "figures/q3/fig14_q3_heaf_framework.pdf"


def main() -> None:
    deck = Presentation(str(SOURCE))
    if len(deck.slides) != 1:
        raise ValueError("Expected one manuscript framework slide")
    seen = set()
    for shape in deck.slides[0].shapes:
        if not shape.has_text_frame:
            continue
        if shape.text in ("音频/视觉证据", "音频/视觉特征"):
            runs = shape.text_frame.paragraphs[0].runs
            assert runs[-1].text in ("证据", "特征")
            runs[-1].text = "特征"
            seen.add("branch")
        elif shape.text in ("仅定位特征槽位", "未对齐特征行"):
            shape.text_frame.paragraphs[0].runs[0].text = "未对齐特征行"
            seen.add("row")
        elif shape.text in ("未核验", "已核验"):
            shape.text_frame.paragraphs[0].runs[0].text = "已核验"
            seen.add("status")
    if seen != {"branch", "row", "status"}:
        raise RuntimeError(f"Figure 14 labels not found: {seen}")
    deck.save(str(SOURCE))

    app = win32.DispatchEx("PowerPoint.Application")
    presentation = None
    try:
        presentation = app.Presentations.Open(str(SOURCE), ReadOnly=True, WithWindow=False)
        presentation.SaveAs(str(PDF), 32)
    finally:
        if presentation is not None:
            presentation.Close()
        app.Quit()
    print(f"Updated manuscript Figure 14: {PDF}")


if __name__ == "__main__":
    main()
