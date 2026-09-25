"""Adapt a copy of the Q2 architecture slide into the Q3 Figure 7 overview.

The retained template geometry is the four dashed outer containers. All Q2
formulae, module contents, and arrows are removed from the copy before the Q3
content is added. Source media are a single paired Attachment4 sample (09).
"""

from __future__ import annotations

import hashlib
import json
import math
import pickle
import shutil
import subprocess
import wave
from pathlib import Path

import cv2
import imageio_ffmpeg
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


E_ROOT = Path(__file__).resolve().parents[5]
OUT = E_ROOT / "outputs" / "final" / "q3" / "figures"
DATA = OUT / "data"
SOURCE = Path(r"C:\Users\17299\Desktop\q2 架构图.pptx")
PPTX = OUT / "figure7_q3_framework_zh_v2.pptx"
SAMPLE_ID = "09"
FRAME = DATA / "figure7_sample09_context_frame.png"
WAV = DATA / "figure7_sample09_audio.wav"
MANIFEST = OUT / "figure7_q3_framework_zh_v2_sources.json"
README = OUT / "figure7_q3_framework_zh_v2_README.md"

BLUE = "A9D3E9"       # text
YELLOW = "EFD98B"     # audio
PINK = "EFB7B2"       # vision
FLOW = "2F6FAE"       # principal arrows
ORANGE = "E6A04B"
INK = "343A40"
MUTED = "65717B"
GRID = "D4DCE2"
PALE_BLUE = "EFF7FB"
PALE_YELLOW = "FCF8EA"
PALE_PINK = "FCF1F0"
PALE_PURPLE = "F6F3FA"
PALE_GREEN = "EAF4ED"
GREEN = "427454"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sample_sources() -> tuple[Path, Path, str, dict]:
    inventory = json.loads((E_ROOT / "data/manifests/q3/attachment4_inventory.json").read_text(encoding="utf-8"))
    entries = inventory["files"]
    video = Path(next(entry["absolute_path"] for entry in entries
                      if entry["absolute_path"].replace("\\", "/").endswith(f"/对齐版本/videos/{SAMPLE_ID}.mp4")))
    pkl = Path(next(entry["absolute_path"] for entry in entries
                    if entry["absolute_path"].replace("\\", "/").endswith(f"/对齐版本/{SAMPLE_ID}.pkl")))
    raw = str(pickle.load(pkl.open("rb"))["raw_text"])
    fragment = "I did not like this movie at all"
    if fragment not in raw:
        raise ValueError("The displayed English fragment is not an exact source substring")
    return video, pkl, fragment, {"raw_text": raw, "fragment_start": raw.index(fragment),
                                  "fragment_end": raw.index(fragment) + len(fragment)}


def create_media(video: Path) -> dict:
    DATA.mkdir(parents=True, exist_ok=True)
    cap = cv2.VideoCapture(str(video))
    if not cap.isOpened():
        raise RuntimeError(f"Cannot open source video: {video}")
    try:
        count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = float(cap.get(cv2.CAP_PROP_FPS))
        decoded = []
        while True:
            ok, bgr = cap.read()
            if not ok:
                break
            decoded.append(bgr)
        if not decoded:
            raise RuntimeError("Source video has no decodable frame")
        actual_index = len(decoded) // 2
        bgr = decoded[actual_index]
        encoded, buffer = cv2.imencode(".png", bgr)
        if not encoded:
            raise RuntimeError("Fixed middle-of-decodable-frames encoding failed")
        buffer.tofile(str(FRAME))  # OpenCV imwrite fails on some Unicode Windows paths.
    finally:
        cap.release()

    command = [imageio_ffmpeg.get_ffmpeg_exe(), "-y", "-loglevel", "error", "-i", str(video),
               "-vn", "-ac", "1", "-ar", "16000", "-sample_fmt", "s16", str(WAV)]
    run = subprocess.run(command, capture_output=True, text=True, check=False)
    if run.returncode:
        raise RuntimeError(f"Audio extraction failed: {run.stderr[-500:]}")
    with wave.open(str(WAV), "rb") as audio:
        if audio.getnchannels() != 1 or audio.getframerate() != 16000 or audio.getsampwidth() != 2:
            raise ValueError("Unexpected WAV format")
        data = np.frombuffer(audio.readframes(audio.getnframes()), dtype="<i2").astype(np.float64) / 32768.0
        audio_length = audio.getnframes() / audio.getframerate()
    if not len(data) or not np.isfinite(data).all() or np.max(np.abs(data)) <= 0:
        raise ValueError("Source audio is empty, silent or nonfinite")
    return {"reported_video_frame_count": count, "decoded_video_frame_count": len(decoded),
            "reported_video_fps": fps,
            "requested_relative_position": 0.5, "decoded_frame_index_zero_based": actual_index,
            "nominal_frame_time_seconds": actual_index / fps,
            "audio_sample_rate_hz": 16000, "audio_duration_seconds": audio_length,
            "audio_samples": data}


def shape(slide, x, y, w, h, *, fill=WHITE, line=GRID, radius=True, width=1.0, dash=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.line.color.rgb = rgb(line)
    item.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if radius:
        try:
            item.adjustments[0] = 0.12
        except (AttributeError, IndexError):
            pass
    return item


def text(slide, content, x, y, w, h, *, size=15, bold=False, color=INK,
         align="center", valign="middle", fill=None):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = rgb(fill)
    tf = item.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.005)
    tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "top": MSO_ANCHOR.TOP,
                          "bottom": MSO_ANCHOR.BOTTOM}[valign]
    for idx, line in enumerate(content.split("\n")):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.alignment = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT,
                               "right": PP_ALIGN.RIGHT}[align]
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return item


def line(slide, x1, y1, x2, y2, *, color=GRID, width=1.0):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    item.line.color.rgb = rgb(color)
    item.line.width = Pt(width)
    return item


def arrow(slide, x1, y1, x2, y2, *, color=FLOW, width=1.5):
    item = line(slide, x1, y1, x2, y2, color=color, width=width)
    tip = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                 Inches(x2-0.075), Inches(y2-0.075), Inches(0.15), Inches(0.15))
    tip.rotation = 90 + math.degrees(math.atan2(y2-y1, x2-x1))
    tip.fill.solid()
    tip.fill.fore_color.rgb = rgb(color)
    tip.line.fill.background()
    return item


def labeled_box(slide, content, x, y, w, h, *, fill=WHITE, stroke=GRID,
                size=14, bold=False, color=INK):
    shape(slide, x, y, w, h, fill=fill, line=stroke, radius=True)
    text(slide, content, x+0.05, y+0.03, w-0.10, h-0.06,
         size=size, bold=bold, color=color)


def retain_framework(slide) -> None:
    group = next(item for item in slide.shapes if item.shape_type == MSO_SHAPE_TYPE.GROUP)
    original = list(group.shapes)
    keep = {0, 2, 4, 6}  # Q2's left, central, upper-right and lower-right containers
    for index, item in enumerate(original):
        if index not in keep:
            group.shapes._spTree.remove(item._element)
    for item in list(slide.shapes):
        if item is not group:
            slide.shapes._spTree.remove(item._element)
    for item, color in zip(list(group.shapes), (BLUE, "D7CDE4", YELLOW, "C9DAD1")):
        item.fill.background()
        item.line.color.rgb = rgb(color)
        item.line.width = Pt(1.1)


def draw_framework_outlines(slide) -> None:
    """Restore the template's visible four-container geometry after regrouping."""
    for x, y, w, h, color in (
        (0.13, 0.82, 3.75, 5.72, BLUE),
        (3.98, 0.68, 5.94, 5.76, "D7CDE4"),
        (10.23, 0.64, 2.85, 3.08, YELLOW),
        (10.23, 3.94, 3.00, 2.57, "C9DAD1"),
    ):
        border = shape(slide, x, y, w, h, fill=WHITE, line=color,
                       radius=True, width=1.15, dash=True)
        border.fill.background()


def waveform(slide, samples: np.ndarray) -> None:
    bins = np.array_split(samples, 70)
    rms = np.array([math.sqrt(float(np.mean(bin_values ** 2))) for bin_values in bins])
    scale = float(np.max(rms))
    line(slide, 0.48, 3.49, 1.94, 3.49, color="D4B875", width=0.45)
    for index, value in enumerate(rms):
        x = 0.48 + 1.46 * index / (len(rms)-1)
        half = 0.025 + 0.31 * value / scale
        line(slide, x, 3.49-half, x, 3.49+half, color="A6772C", width=1.0)


def draw_left(slide, fragment: str, samples: np.ndarray) -> None:
    labeled_box(slide, "①  多模态输入与表示", 0.85, 0.53, 2.34, 0.43,
                fill=WHITE, stroke=BLUE, size=16, bold=True)
    for y in (2.78, 4.48):
        line(slide, 0.35, y, 3.72, y, color="E2E8EC", width=0.8)
    for y, title, color in ((1.20,"文本",BLUE),(2.95,"音频",YELLOW),(4.63,"视觉",PINK)):
        shape(slide, 0.36, y+0.02, 0.10, 0.28, fill=color, line=color, radius=False)
        text(slide, title, 0.52, y, 0.68, 0.34, size=16, bold=True, align="left")
    labeled_box(slide, f'“{fragment}”', 0.39, 1.57, 1.64, 0.69,
                fill=PALE_BLUE, stroke=BLUE, size=11.2)
    text(slide, "原文片段", 0.43, 2.29, 1.56, 0.22, size=10.5, color=MUTED)
    waveform(slide, samples)
    text(slide, "原视频音轨波形", 0.40, 3.87, 1.64, 0.24, size=10.5, color=MUTED)
    slide.shapes.add_picture(str(FRAME), Inches(0.47), Inches(4.99), width=Inches(1.50))
    text(slide, "原视频场景（非定位帧）", 0.34, 6.16, 1.82, 0.21,
         size=9.6, color=MUTED)
    features = ((1.72, "文本特征", "50 × 768", PALE_BLUE, BLUE),
                (3.42, "音频特征", "50 × 74", PALE_YELLOW, YELLOW),
                (5.13, "视觉特征", "50 × 35", PALE_PINK, PINK))
    for y, title, dims, fill, stroke in features:
        arrow(slide, 2.06, y+0.34, 2.26, y+0.34, width=1.4)
        labeled_box(slide, f"{title}\n{dims}", 2.31, y, 1.34, 0.71,
                    fill=fill, stroke=stroke, size=13.5, bold=True)


def draw_predictor(slide) -> None:
    labeled_box(slide, "②  情感预测与 HEAF 解释", 5.14, 0.46, 3.58, 0.44,
                fill=WHITE, stroke="D7CDE4", size=16, bold=True)
    text(slide, "情感预测", 4.25, 1.06, 1.30, 0.31, size=16, bold=True,
         align="left", color=FLOW)
    labeled_box(slide, "三模态\n特征", 4.25, 1.51, 1.07, 0.85,
                fill=WHITE, stroke=GRID, size=14, bold=True)
    for i, color in enumerate((BLUE,YELLOW,PINK)):
        shape(slide, 4.43+i*0.23, 2.20, 0.17, 0.08,
              fill=color, line=color, radius=False)
    arrow(slide, 5.35, 1.93, 5.60, 1.93)
    labeled_box(slide, "均值池化\n＋注意力残差", 5.63, 1.51, 1.45, 0.85,
                fill=PALE_BLUE, stroke=BLUE, size=13, bold=True)
    arrow(slide, 7.10, 1.93, 7.35, 1.93)
    labeled_box(slide, "多模态\n融合", 7.38, 1.51, 0.98, 0.85,
                fill=WHITE, stroke=GRID, size=13.5, bold=True)
    arrow(slide, 8.39, 1.93, 8.58, 1.93)
    labeled_box(slide, "类别预测\n强度预测", 8.61, 1.51, 1.02, 0.85,
                fill=PALE_PURPLE, stroke="C9BCD8", size=13.5, bold=True)
    line(slide, 4.20, 2.71, 9.67, 2.71, color=GRID, width=0.8)
    arrow(slide, 6.92, 2.41, 6.92, 3.00, width=1.25)
    text(slide, "HEAF 层级解释", 4.25, 2.99, 2.52, 0.35,
         size=16, bold=True, align="left", color=FLOW)


def draw_heaf(slide) -> None:
    card_x = (4.25, 5.64, 7.03, 8.42)
    titles = ("八种联盟", "Shapley贡献", "两两交互", "时间遮挡")
    for x, title in zip(card_x, titles):
        shape(slide, x, 3.53, 1.26, 1.55, fill=WHITE, line=GRID)
        text(slide, title, x+0.04, 4.66, 1.18, 0.30, size=12.3, bold=True)
    for i in range(3):
        arrow(slide, card_x[i]+1.27, 4.28, card_x[i+1]-0.03, 4.28, width=1.25)
    # 8 coalition chips; equal size encodes enumeration, not a score.
    for i in range(8):
        x = 4.37 + (i % 4)*0.28
        y = 3.83 + (i // 4)*0.30
        shape(slide, x, y, 0.22, 0.20, fill=PALE_BLUE, line=BLUE,
              radius=False, width=0.7)
    # Three bars identify three modalities, without implying numeric values.
    for i, (col, w) in enumerate(((BLUE,0.74),(YELLOW,0.62),(PINK,0.52))):
        shape(slide, 5.83, 3.80+i*0.27, w, 0.13, fill=col, line=col,
              radius=False, width=0.3)
    for x, y, col in ((7.32,3.89,BLUE),(7.72,3.89,YELLOW),(7.52,4.30,PINK)):
        shape(slide, x, y, 0.13, 0.13, fill=col, line=col, radius=True)
    line(slide, 7.40, 3.97, 7.78, 3.97, color=MUTED, width=1.0)
    line(slide, 7.40, 4.00, 7.56, 4.30, color=MUTED, width=1.0)
    line(slide, 7.78, 4.00, 7.57, 4.30, color=MUTED, width=1.0)
    for i in range(10):
        col = "FBE5C7" if 4 <= i <= 6 else "E7EEF2"
        shape(slide, 8.52+i*0.105, 4.00, 0.09, 0.32,
              fill=col, line=WHITE, radius=False, width=0.1)
    text(slide, "连续特征槽位", 8.48, 4.35, 1.13, 0.20, size=9.5, color=MUTED)
    labeled_box(slide, "主导模态  ·  模态交互  ·  关键特征区间", 4.36, 5.47, 5.19, 0.50,
                fill=PALE_PURPLE, stroke="D7CDE4", size=14, bold=True)
    text(slide, "精确归因  →  局部遮挡  →  删除对照", 4.50, 6.06, 4.93, 0.21,
         size=10.7, color=MUTED)


def draw_right(slide) -> None:
    labeled_box(slide, "③  解释结果输出", 10.58, 0.48, 2.14, 0.43,
                fill=WHITE, stroke=YELLOW, size=16, bold=True)
    labeled_box(slide, "解释卡", 10.61, 1.13, 2.07, 0.39,
                fill=PALE_YELLOW, stroke=YELLOW, size=15, bold=True)
    for i, (title, value) in enumerate((
        ("预测", "类别 · 情感强度"),
        ("模态", "贡献 · 主导模态"),
        ("局部", "交互 · 特征区间"),
    )):
        y = 1.64+i*0.58
        labeled_box(slide, title, 10.47, y, 0.64, 0.43,
                    fill=WHITE, stroke=GRID, size=12.3, bold=True)
        labeled_box(slide, value, 11.22, y, 1.52, 0.43,
                    fill=WHITE, stroke=GRID, size=11.3)
    arrow(slide, 11.67, 3.75, 11.67, 4.06, width=1.7)
    labeled_box(slide, "证据回溯与有效性", 10.50, 3.80, 2.42, 0.43,
                fill=WHITE, stroke="C9DAD1", size=15.2, bold=True)
    labeled_box(slide, "文本证据：可回溯", 10.49, 4.45, 2.47, 0.42,
                fill=PALE_GREEN, stroke="BFD7C7", size=12.4,
                bold=True, color=GREEN)
    labeled_box(slide, "音 / 视：仅特征槽位", 10.49, 4.99, 2.47, 0.42,
                fill="F5F6F7", stroke=GRID, size=12.2, color=MUTED)
    text(slide, "原始秒数与帧号未验证", 10.48, 5.41, 2.48, 0.25,
         size=10.5, color=MUTED)
    labeled_box(slide, "关键区间删除  ↔  随机删除", 10.49, 5.82, 2.47, 0.44,
                fill=PALE_BLUE, stroke=BLUE, size=11.7)


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    video, pkl, fragment, text_meta = sample_sources()
    media = create_media(video)
    shutil.copy2(SOURCE, PPTX)
    presentation = Presentation(str(PPTX))
    if len(presentation.slides) != 1:
        raise ValueError("Expected the one-slide Q2 architecture source")
    slide = presentation.slides[0]
    retain_framework(slide)
    draw_framework_outlines(slide)
    text(slide, "图7  多模态情感预测与层级证据归因框架", 0.6, 0.04, 12.15, 0.38,
         size=20, bold=True)
    draw_left(slide, fragment, media["audio_samples"])
    arrow(slide, 3.78, 3.49, 4.17, 3.49, width=2.0)
    draw_predictor(slide)
    draw_heaf(slide)
    arrow(slide, 9.90, 1.94, 10.28, 1.94, width=1.8)
    arrow(slide, 9.90, 5.34, 10.28, 5.34, width=1.8)
    draw_right(slide)
    text(slide, "原视频画面与音频波形仅作输入示例，不表示解释模型定位的关键时间或帧。",
         0.55, 6.76, 12.21, 0.25, size=11.0, color=MUTED)
    presentation.save(str(PPTX))

    reopened = Presentation(str(PPTX))
    if len(reopened.slides) != 1 or len(reopened.slides[0].shapes) < 100:
        raise RuntimeError("Saved Figure 7 PPTX failed basic integrity checks")
    manifest = {
        "figure": PPTX.name,
        "template_pptx": str(SOURCE), "template_sha256": sha256(SOURCE),
        "template_framework_retained": "four original dashed outer containers and slide geometry",
        "sample_id": SAMPLE_ID, "selection_basis": "new original scene and concise sentiment sentence; independent of HEAF outputs",
        "feature_pkl": str(pkl), "feature_pkl_sha256": sha256(pkl),
        "raw_text": text_meta["raw_text"],
        "displayed_fragment": fragment,
        "displayed_fragment_character_span": [text_meta["fragment_start"], text_meta["fragment_end"]],
        "source_mp4": str(video), "source_mp4_sha256": sha256(video),
        "frame_file": str(FRAME), "frame_sha256": sha256(FRAME),
        "frame_selection_method": "middle frame of sequentially decodable video frames, no HEAF information",
        "requested_relative_position": media["requested_relative_position"],
        "reported_video_frame_count": media["reported_video_frame_count"],
        "decoded_video_frame_count": media["decoded_video_frame_count"],
        "actual_decoded_frame_index_zero_based": media["decoded_frame_index_zero_based"],
        "nominal_decoded_time_seconds": media["nominal_frame_time_seconds"],
        "frame_purpose": "original_video_context_only",
        "frame_grounding_status": "not_keyframe_mapping",
        "audio_wav": str(WAV), "audio_wav_sha256": sha256(WAV),
        "audio_extraction": "paired MP4 full audio; mono PCM16 16 kHz",
        "audio_sample_rate_hz": media["audio_sample_rate_hz"],
        "audio_duration_seconds": media["audio_duration_seconds"],
        "waveform_display": "70 equal-duration RMS bins scaled by the sample maximum; no feature-slot alignment",
        "raw_audio_visual_grounding": "unverified",
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    README.write_text(
        "# Figure 7（Q3）新版说明\n\n"
        "本图从 `q2 架构图.pptx` 复制一份作为可编辑底座，保留原图四个虚线分区的几何布局，"
        "删除其中 Q2 公式、旧箭头和旧模块，再填入 Q3 的输入、情感预测、HEAF 分析与解释输出。"
        "原始 Q2 PPTX 和已有 Figure 7 文件均未修改。\n\n"
        "左侧三个真实素材均来自附件4配对样本09：逐字原文片段、可解码视频序列的中间帧原视频场景、"
        "以及该 MP4 全音轨提取后计算的真实 RMS 波形。媒体素材只表示输入；视频帧不是关键帧，"
        "音频波形不对应 HEAF 的特征槽位。文本原文保留英文，其余图中文字为中文。"
        "来源哈希与选择位置见 `figure7_q3_framework_zh_v2_sources.json`。\n\n"
        "图注建议：图7 多模态情感预测与层级证据归因框架。三模态特征经池化与融合形成情感类别和"
        "强度预测，HEAF 再进行精确模态归因、交互分析与连续窗口遮挡，输出解释卡及有效性对照。"
        "左侧视频画面和音频波形为同一原始样本的输入示例，不表示模型定位的关键帧或关键音频时间；"
        "文本证据可回溯原文，音视频原始秒数与帧号尚未验证。\n",
        encoding="utf-8",
    )
    print(f"CREATED {PPTX} shapes={len(reopened.slides[0].shapes)}")


if __name__ == "__main__":
    build()
