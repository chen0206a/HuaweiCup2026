"""Read-only evidence audit. Run from repository root; never load attachment3.

Recomputes saved metrics, checks source-video splits and checkpoint identity.
Does not train, select a new model, or evaluate held-out examples.
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import statistics as st

from openpyxl import load_workbook
import torch

ROOT = Path(__file__).resolve().parents[1]
Q1 = ROOT / "Q1_方法流程与实验结果_交付包"
Q2 = ROOT / "E2026"


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def score(row):
    return .25 * (row["accuracy"] + row["macro_f1"] + 1 - row["mae"] / 6 + (row["pearson"] + 1) / 2)


def main():
    out = {"scope": "saved results, code contracts, checkpoint tensors and ID metadata only; no training or dataset inference"}
    splits = read(Q1 / "outputs/q1_visual_encoder_ablation/cv_splits.json")["splits"]
    overlaps = []
    for s in splits:
        tr = {x.rsplit("__", 1)[0] for x in s["train_sample_ids"]}
        te = {x.rsplit("__", 1)[0] for x in s["test_sample_ids"]}
        overlaps.append({"repeat": s["repeat"], "fold": s["fold"], "shared_video_count": len(tr & te),
                         "test_clips_with_train_video": sum(x.rsplit("__", 1)[0] in tr for x in s["test_sample_ids"])})
    out["q1_cv_group_overlap"] = overlaps
    with (Q1 / "outputs/q1_alignment_ablation/downstream_cv_results.csv").open(encoding="utf-8-sig", newline="") as f:
        raw = list(csv.DictReader(f))
    with (Q1 / "outputs/q1_alignment_ablation/paper_alignment_comparison.csv").open(encoding="utf-8-sig", newline="") as f:
        reported = list(csv.DictReader(f))
    errors = []
    for row in reported:
        for metric in ("accuracy", "macro_f1", "mae", "pearson"):
            vals = [float(x[metric]) for x in raw if x["configuration"] == row["configuration"]]
            assert len(vals) == 25
            errors.extend([abs(st.mean(vals) - float(row[metric + "_mean"])), abs(st.stdev(vals) - float(row[metric + "_std"]))])
    out["q1_alignment_table_max_abs_recompute_error"] = max(errors)
    expected = ["outputs/q1_cpu/samples", "outputs/q1_features_100_final", "outputs/q1_text_final/text_feature_sources/text_feature_sources.json",
                "outputs/q1_final_local/01_final_features", "outputs/q1_final_local/05_appendix", "scripts/build_q1_summary_xlsx.mjs"]
    out["q1_delivery_required_dependencies_exist"] = {p: (Q1 / p).exists() for p in expected}

    # Inspect identifiers only. No test features or predictions are evaluated.
    wb = load_workbook(Q2 / "data/raw/label.xlsx", read_only=True, data_only=True)
    it = wb.active.iter_rows(values_only=True)
    header = next(it)
    groups = {}
    counts = {}
    for row in it:
        d = dict(zip(header, row))
        groups.setdefault(d["mode"], set()).add(str(d["video_id"]))
        counts[d["mode"]] = counts.get(d["mode"], 0) + 1
    wb.close()
    out["q2_split_metadata"] = {"sample_counts": counts, "video_counts": {k: len(v) for k, v in groups.items()},
                                "shared_video_counts": {a + "__" + b: len(groups[a] & groups[b]) for a in groups for b in groups if a < b}}
    lock = read(Q2 / "outputs/final/q2/q2_model_lock.json")
    out["q2_hash_checks"] = {k: sha(Q2 / v["relative_path"]) == v["sha256"] for k, v in lock["checkpoint_paths"].items()}
    out["q2_hash_checks"]["benchmark"] = sha(Q2 / "outputs/metrics/b2_benchmark_definition.json") == lock["benchmark"]["sha256"]
    out["q2_frozen_checkpoint_tensors"] = {}
    for seed in (42, 43, 44):
        b = torch.load(Q2 / lock["checkpoint_paths"][f"B0-WCE:{seed}"]["relative_path"], map_location="cpu", weights_only=False)["model_state_dict"]
        p = torch.load(Q2 / lock["checkpoint_paths"][f"B5-P2:{seed}"]["relative_path"], map_location="cpu", weights_only=False)["model_state_dict"]
        changed = [k for k in b if not torch.equal(b[k], p[k])]
        out["q2_frozen_checkpoint_tensors"][str(seed)] = {"b0_tensor_count": len(b), "changed": changed}
    d = read(Q2 / "outputs/metrics/b5_p2_multiseed_summary.json")
    screen = read(Q2 / "outputs/metrics/b5_pooling_metrics.json")
    errors = []
    for seed, detail in d["per_seed"].items():
        for model in ("B0", "P2"):
            val = detail[model]
            for condition in ("clean", "mean_missing"):
                errors.append(abs(score(val[condition]) - val[condition]["selection_score"]))
            errors.append(abs(.5 * (val["clean"]["selection_score"] + val["mean_missing"]["selection_score"]) - val["robust_score"]))
        for model, aggregate_name in (("B0", "B0-WCE"), ("P2", "B5-P2")):
            for condition in ("clean", "mean_missing"):
                for metric in ("accuracy", "macro_f1", "mae", "pearson", "selection_score"):
                    vals = [v[model][condition][metric] for v in d["per_seed"].values()]
                    agg = d["models"][aggregate_name][condition][metric]
                    errors.extend([abs(st.mean(vals) - agg["mean"]), abs(st.stdev(vals) - agg["std"])])
        rows = screen["P2"]["scenario_details"] if seed == "42" else d["seed43_44_details"][seed]["best_robust_scenario_details"]
        assert len(rows) == 55 and rows[0]["scenario_id"] == "clean" and len({r["scenario_id"] for r in rows}) == 55
        for metric in ("accuracy", "macro_f1", "mae", "pearson", "selection_score"):
            errors.append(abs(st.mean(r[metric] for r in rows[1:]) - detail["P2"]["mean_missing"][metric]))
    out["q2_saved_metric_max_abs_recompute_error"] = max(errors)
    deltas = [v["P2"]["robust_score"] - v["B0"]["robust_score"] for v in d["per_seed"].values()]
    clean_deltas = []
    for seed in ("42", "43", "44"):
        v = d["seed42_best_clean_checkpoint"]["validation"] if seed == "42" else d["seed43_44_details"][seed]["best_clean_validation"]
        clean_deltas.append(v["robust_score"] - d["per_seed"][seed]["B0"]["robust_score"])
    out["q2_paired_robust_delta"] = {"values": deltas, "mean": st.mean(deltas), "sample_sd": st.stdev(deltas)}
    out["q2_both_clean_selected_paired_robust_delta"] = {"values": clean_deltas, "mean": st.mean(clean_deltas), "sample_sd": st.stdev(clean_deltas)}
    out["q2_clean_to_missing_score_drop"] = {k: v["clean"]["selection_score"]["mean"] - v["mean_missing"]["selection_score"]["mean"] for k, v in d["models"].items()}
    out["source_hashes"] = {str(p.relative_to(ROOT)): sha(p) for p in [Q2 / "outputs/metrics/b5_p2_multiseed_summary.json", Q2 / "outputs/metrics/b5_pooling_metrics.json", Q1 / "outputs/q1_visual_encoder_ablation/cv_splits.json", Q1 / "outputs/q1_final_local/00_manifest/q1_paper_numbers.json"]}
    assert all(out["q2_hash_checks"].values())
    assert not any(x["changed"] for x in out["q2_frozen_checkpoint_tensors"].values())
    assert out["q1_alignment_table_max_abs_recompute_error"] < 1e-12 and max(errors) < 1e-12
    dest = ROOT / "notes/e_q1_q2_audit_evidence_20260924.json"
    dest.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(dest)
    print("Q1/Q2 saved-number recomputation and six checkpoint checks passed.")


if __name__ == "__main__":
    main()
