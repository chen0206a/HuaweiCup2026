from pathlib import Path
import re,json,hashlib,csv
import pymupdf as fitz
from pptx import Presentation
from zipfile import ZipFile
from PIL import Image,ImageOps,ImageDraw
O=Path(__file__).resolve().parent
ROOT=O.parent.parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
tex=(O/'q3_story_refine.tex').read_text(encoding='utf-8')
source=(O/'source/07_q3.tex').read_text(encoding='utf-8-sig')
checks=[]
def check(name,ok,detail=''):
    checks.append({'name':name,'pass':bool(ok),'detail':detail})
for e in re.findall(r'\\begin\{equation\}.*?\\end\{equation\}',source,re.S):
    check('公式原样保留 '+re.search(r'\\label\{([^}]+)\}',e)[1],e in tex)
for old in re.findall(r'\\begin\{table\}.*?\\end\{table\}',source,re.S):
    k=re.search(r'\\label\{([^}]+)\}',old)[1]
    new=next(t for t in re.findall(r'\\begin\{table\}.*?\\end\{table\}',tex,re.S) if '\\label{'+k+'}' in t)
    # Numeric body content: caption changes cannot change results.
    old=old.split(r'\midrule',1)[1].split(r'\bottomrule')[0]
    new=new.split(r'\midrule',1)[1].split(r'\bottomrule')[0]
    check('结果表数字不变 '+k,re.findall(r'[+-]?\d+(?:\.\d+)?',old)==re.findall(r'[+-]?\d+(?:\.\d+)?',new))
PAPER=ROOT/'论文润色工作区/E2026_论文整理包/01_LaTeX论文工程/paper'
for name in ('table_q3_attachment4_all.tex','table_q3_attachment4_explanations.tex'):
    old=(PAPER/'generated'/name).read_text(encoding='utf-8-sig')
    new=(O/'generated'/name).read_text(encoding='utf-8')
    check('20条完整保留 '+name,len(re.findall(r'\\detokenize\{\d{2}\}',new))==20)
    check('明细全部数字不变 '+name,re.findall(r'[+-]?\d+(?:\.\d+)?',old)==re.findall(r'[+-]?\d+(?:\.\d+)?',new))
csvpath=ROOT/'E2026/outputs/q3/final/attachment4_predictions_explanations.csv'
rows={r['sample_id']:r for r in csv.DictReader(csvpath.open(encoding='utf-8-sig',newline=''))}
classlabels={'Negative':'消极','Neutral':'中性','Positive':'积极'}
for name in ('table_q3_attachment4_all.tex','table_q3_attachment4_explanations.tex'):
    correct=True
    for line in (O/'generated'/name).read_text(encoding='utf-8').splitlines():
        found=re.search(r'\\detokenize\{(\d{2})\}',line)
        if not found:continue
        row=rows[found[1]]
        parts=[p.strip().removesuffix('\\\\') for p in line.split('&')]
        if name.endswith('_all.tex'):
            correct &= parts[1]==classlabels[row['predicted_class_name']]
            correct &= parts[2]==f"{float(row['predicted_intensity']):+.6f}"
            correct &= parts[3].strip()==f"{float(row['confidence']):.4f}"
        else:
            correct &= all(parts[i+1]==f"{float(row['shapley_'+mod]):+.4f}" for i,mod in enumerate(('text','audio','vision')))
            correct &= parts[4]=={'text':'文本','audio':'语音','vision':'视觉'}[row['primary_modality_classification']]
            correct &= parts[5]==f"$[{row['key_start_index']},{row['key_end_index']})$"
    check('直接核对冻结CSV '+name,correct)
protected=json.loads((O/'qa/protected_hashes.json').read_text(encoding='utf-8'))
for path,h in protected.items():check('受保护文件 '+path,sha(Path(path))==h)
for term in ('层级','多层次','证据链','HEAF','对数赔率','关键时间','方法选择依据','\\textbf'):
    text=re.sub(r'\\(?:label|ref)\{[^}]+\}|figures/[^}]+','',tex)
    check('正文移除 '+term,term not in text)
for term in ('Spearman','0.400735','0.103545','0.108296','0.237758','0.418076','0.465465','0.869','0.831','0.579','0.527','604','564/564','534/534'):
    check('关键事实 '+term,term in tex)
check('8个二级节',len(re.findall(r'\\subsection\{',tex))==8)
refs=re.findall(r'\\(?:ref|eqref)\{([^}]+)\}',tex)
labels=re.findall(r'\\label\{([^}]+)\}',tex)
for p in (O/'generated').glob('*.tex'):labels+=re.findall(r'\\label\{([^}]+)\}',p.read_text(encoding='utf-8'))
check('引用全部有定义',all(r in labels for r in refs))
log=(O/'build/q3_story_refine.log').read_text(encoding='utf-8',errors='replace')
check('无未定义引用',not re.search('undefined|Undefined',log))
check('无overfull',not re.search('Overfull',log))
check('无缺字',not re.search('Missing character',log))
figdir=O/'figures/q3'
figtext='\n'.join(fitz.open(q)[0].get_text() for q in figdir.glob('*.pdf'))
for term in ('HEAF','已核验','未核验','对数赔率','分类边际','类别边际','负面','正面','关键时间'):
    check('图件术语清理 '+term,term not in figtext)
oldppt=ROOT/'论文整理/paper_revision_work/paper/figures/q3/fig14_q3_heaf_framework_source.pptx'
newppt=O/'figure_sources/fig14_q3_framework_refined_source.pptx'
a,b=Presentation(oldppt),Presentation(newppt)
check('流程图重开一页',len(b.slides)==len(a.slides)==1)
check('流程图所有形状几何不变',[(s.shape_id,s.left,s.top,s.width,s.height) for s in a.slides[0].shapes]==[(s.shape_id,s.left,s.top,s.width,s.height) for s in b.slides[0].shapes])
check('流程图全部文本数字不变', sorted(re.findall(r'\d+',' '.join(s.text for s in a.slides[0].shapes if s.has_text_frame)))==sorted(re.findall(r'\d+',' '.join(s.text for s in b.slides[0].shapes if s.has_text_frame))))
with ZipFile(oldppt) as za,ZipFile(newppt) as zb:
    media=[n for n in za.namelist() if n.startswith('ppt/media/')]
    check('流程图媒体全部不变',all(za.read(n)==zb.read(n) for n in media))
p=fitz.open(O/'q3_story_refine.pdf')
check('页数12至14',12<=len(p)<=14,str(len(p)))
text='\n'.join(page.get_text() for page in p)
(O/'qa/pdf_text.txt').write_text(text,encoding='utf-8')
for i,page in enumerate(p):
    pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
    pix.save(O/'qa'/f'page-{i+1:02}.png')
# 3-page sheets: every final page must be visually inspected.
for start in range(0,len(p),3):
    thumbs=[]
    for index in range(start,min(start+3,len(p))):
        img=Image.open(O/'qa'/f'page-{index+1:02}.png').convert('RGB')
        img.thumbnail((690,976))
        tile=Image.new('RGB',(710,1010),'#ddd')
        tile.paste(img,((710-img.width)//2,28))
        ImageDraw.Draw(tile).text((12,8),f'Page {index+1}',fill='black')
        thumbs.append(tile)
    canvas=Image.new('RGB',(710*len(thumbs),1010),'white')
    for index,img in enumerate(thumbs):canvas.paste(img,(index*710,0))
    canvas.save(O/'qa'/f'sheet-{start//3+1}.png')
(O/'qa/checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'pages':len(p),'checks':len(checks),'passed':sum(x['pass'] for x in checks),'failures':[x for x in checks if not x['pass']]},ensure_ascii=False))
